#!/usr/bin/env python3
"""Extrai grupos e metadados de planilha preenchida para JSON."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import openpyxl

from ifpug_calc import calcular_pf

COL_FUNCAO = 3
COL_TIPO = 8
COL_TD = 10
COL_AR = 11
COL_IAET = 9
COL_PF = 15
COL_PF_LOCAL = 16
COL_OBS = 17
LINHA_INICIO = 12

FATORES_IAET = {"I": "W15", "A": "W16", "E": "W17", "T": "W18"}


def fator_iaet(ws_contagem, iaet: str | None) -> float:
    iaet = (iaet or "I").upper()
    celula = FATORES_IAET.get(iaet, "W15")
    valor = ws_contagem[celula].value
    return float(valor) if valor is not None else 1.0


def extrair(caminho: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws_c = wb["Contagem"]
    ws_f = wb["Funções"]

    identificacao = {
        "empresa": ws_c["F7"].value,
        "aplicacao": ws_c["F8"].value,
        "projeto": ws_c["F9"].value,
        "dv_numero": ws_c["F10"].value,
        "responsavel": ws_c["F11"].value,
        "versao": ws_c["N10"].value,
        "revisor": ws_c["F12"].value,
        "data_contagem": ws_c["W11"].value,
        "rs_por_pf": ws_c["T7"].value,
        "proposito": ws_c["K20"].value,
    }

    grupos: list[dict[str, Any]] = []
    grupo_atual: dict[str, Any] | None = None
    idx = 0

    for r in range(LINHA_INICIO, ws_f.max_row + 1):
        nome = ws_f.cell(r, COL_FUNCAO).value
        tipo = ws_f.cell(r, COL_TIPO).value
        if not nome:
            continue
        nome = str(nome).strip()
        if not tipo:
            idx += 1
            grupo_atual = {"id": f"F{idx}", "nome": nome, "itens": []}
            grupos.append(grupo_atual)
            continue
        if grupo_atual is None:
            continue
        td = ws_f.cell(r, COL_TD).value
        ar = ws_f.cell(r, COL_AR).value
        iaet = ws_f.cell(r, COL_IAET).value
        pf = ws_f.cell(r, COL_PF).value
        pf_local = ws_f.cell(r, COL_PF_LOCAL).value
        if pf_local is None:
            base = pf if pf is not None else calcular_pf(str(tipo), td, ar)
            pf_local = float(base or 0) * fator_iaet(ws_c, iaet)
        grupo_atual["itens"].append({
            "nome": nome,
            "tipo": str(tipo).strip(),
            "td": td,
            "ar": ar,
            "pf": pf if pf is not None else calcular_pf(str(tipo), td, ar),
            "pf_local": pf_local,
            "iaet": iaet,
            "observacoes": ws_f.cell(r, COL_OBS).value,
        })

    total = sum(
        float(item.get("pf_local") or 0)
        for g in grupos
        for item in g.get("itens", [])
    )

    return {
        "identificacao": identificacao,
        "grupos": grupos,
        "total_pf": total,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Extrai JSON de planilha de contagem")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("-o", "--saida", type=Path, required=True)
    args = parser.parse_args()
    dados = extrair(args.xlsx)
    args.saida.parent.mkdir(parents=True, exist_ok=True)
    with args.saida.open("w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2, default=str)
    print(f"JSON extraído: {args.saida} ({dados['total_pf']} PF)")


if __name__ == "__main__":
    main()
