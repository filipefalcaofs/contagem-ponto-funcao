#!/usr/bin/env python3
"""Preenche planilha de contagem IFPUG a partir de JSON e template Sudoeste."""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path
from typing import Any

import openpyxl

COL_FUNCAO = 3   # C
COL_TIPO = 8     # H
COL_IAET = 9     # I
COL_TD = 10      # J
COL_AR = 11      # K
COL_OBS = 17     # Q
LINHA_INICIO = 12

TIPOS_TRANSACIONAIS = {"EE", "CE", "SE"}
TIPOS_CONTAVEIS = {"EE", "CE", "SE", "ALI", "AIE"}
TERMOS_SUSPEITOS = (
    "gerar pdf",
    "assinar",
    "enfileirar",
    "processar",
    "reprocessar",
    "registrar log",
    "registrar consulta",
    "notificar",
    "enviar e-mail",
    "aplicar template",
    "pré-validar",
    "preencher formulário",
    "calcular hash",
    "vincular retificação",
)
JUSTIFICATIVAS_PE = (
    "processo elementar independente",
    "saída funcional autônoma",
    "acao independente",
    "ação independente",
    "usuário aciona explicitamente",
    "usuario aciona explicitamente",
)


def carregar_json(caminho: Path) -> dict[str, Any]:
    with caminho.open(encoding="utf-8") as f:
        return json.load(f)


def validar_processos_elementares(dados: dict[str, Any]) -> None:
    erros: list[str] = []

    for grupo in dados.get("grupos", []):
        for item in grupo.get("itens", []):
            tipo = item.get("tipo")
            nome = str(item.get("nome", ""))
            observacoes = str(item.get("observacoes", ""))
            texto = f"{nome} {observacoes}".lower()

            if tipo not in TIPOS_CONTAVEIS:
                erros.append(f"{grupo.get('nome')} / {nome}: tipo funcional inválido ({tipo!r}).")
                continue

            if tipo in TIPOS_TRANSACIONAIS:
                if not observacoes.strip():
                    erros.append(f"{grupo.get('nome')} / {nome}: transação sem observações auditáveis.")

                if any(termo in texto for termo in TERMOS_SUSPEITOS) and not any(j in texto for j in JUSTIFICATIVAS_PE):
                    erros.append(
                        f"{grupo.get('nome')} / {nome}: item suspeito de não ser processo elementar independente. "
                        "Incorpore ao processo principal ou justifique explicitamente em observações."
                    )

    if erros:
        print("Validação de processo elementar falhou:", file=sys.stderr)
        for erro in erros:
            print(f"- {erro}", file=sys.stderr)
        raise SystemExit(2)


def iaet_padrao(dados: dict[str, Any], item: dict[str, Any]) -> str | None:
    if item.get("iaet"):
        return item["iaet"]

    tipo_contagem = dados.get("identificacao", {}).get("tipo_contagem", "desenvolvimento").lower()
    if tipo_contagem == "desenvolvimento" and item.get("tipo") in TIPOS_CONTAVEIS:
        return "I"

    return None


def limpar_dados_funcoes(ws, linha_inicio: int = LINHA_INICIO) -> None:
    for r in range(linha_inicio, ws.max_row + 1):
        for col in (COL_FUNCAO, COL_TIPO, COL_IAET, COL_TD, COL_AR, COL_OBS):
            ws.cell(r, col).value = None


def preencher_identificacao(ws, meta: dict[str, Any]) -> None:
    mapeamento = {
        "empresa": ("F", 7),
        "aplicacao": ("F", 8),
        "projeto": ("F", 9),
        "dv_numero": ("F", 10),
        "responsavel": ("F", 11),
        "versao": ("N", 10),
        "revisor": ("F", 12),
        "data_contagem": ("W", 11),
        "rs_por_pf": ("T", 7),
        "proposito": ("K", 20),
    }
    for chave, (col, row) in mapeamento.items():
        if chave in meta and meta[chave] is not None:
            ws[f"{col}{row}"] = meta[chave]

    tipo = meta.get("tipo_contagem", "desenvolvimento").lower()
    linhas_tipo = {
        "estimativa": 14,
        "desenvolvimento": 15,
        "melhoria": 16,
        "baseline": 17,
        "aplicacao": 17,
    }
    for row in linhas_tipo.values():
        ws.cell(row, 12).value = None
    if tipo in linhas_tipo:
        ws.cell(linhas_tipo[tipo], 12).value = "x"


def escrever_grupos(ws, dados: dict[str, Any], linha_inicio: int = LINHA_INICIO) -> int:
    linha = linha_inicio
    for grupo in dados.get("grupos", []):
        ws.cell(linha, COL_FUNCAO).value = grupo["nome"]
        linha += 1
        for item in grupo.get("itens", []):
            ws.cell(linha, COL_FUNCAO).value = item["nome"]
            ws.cell(linha, COL_TIPO).value = item["tipo"]
            ws.cell(linha, COL_IAET).value = iaet_padrao(dados, item)
            if item.get("td") is not None:
                ws.cell(linha, COL_TD).value = item["td"]
            if item.get("ar") is not None:
                ws.cell(linha, COL_AR).value = item["ar"]
            if item.get("observacoes"):
                ws.cell(linha, COL_OBS).value = item["observacoes"]
            linha += 1
        linha += 1  # linha em branco entre grupos
    return linha


def main() -> None:
    parser = argparse.ArgumentParser(description="Preenche planilha de contagem IFPUG")
    parser.add_argument("json", type=Path, help="Arquivo JSON com metadados e grupos")
    parser.add_argument(
        "-o", "--saida", type=Path, required=True, help="Caminho do XLSX de saída"
    )
    parser.add_argument(
        "-t",
        "--template",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "templates" / "planilha-contagem-modelo.xlsx",
        help="Template XLSX",
    )
    args = parser.parse_args()

    dados = carregar_json(args.json)
    validar_processos_elementares(dados)
    args.saida.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.template, args.saida)

    wb = openpyxl.load_workbook(args.saida)
    ws_contagem = wb["Contagem"]
    ws_funcoes = wb["Funções"]

    preencher_identificacao(ws_contagem, dados.get("identificacao", {}))
    limpar_dados_funcoes(ws_funcoes)
    escrever_grupos(ws_funcoes, dados)

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    wb.save(args.saida)
    print(f"Planilha gerada: {args.saida}")


if __name__ == "__main__":
    main()
