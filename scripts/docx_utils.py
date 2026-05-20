"""Utilitários para edição de DOCX preservando formatação do template."""

from __future__ import annotations

from copy import deepcopy
from typing import Any

from docx.table import Table
from docx.text.paragraph import Paragraph


def substituir_texto_documento(doc, substituicoes: dict[str, str]) -> None:
    for paragrafo in doc.paragraphs:
        _substituir_paragrafo(paragrafo, substituicoes)
    for tabela in doc.tables:
        for row in tabela.rows:
            for cell in row.cells:
                for paragrafo in cell.paragraphs:
                    _substituir_paragrafo(paragrafo, substituicoes)
    for secao in doc.sections:
        for paragrafo in secao.header.paragraphs:
            _substituir_paragrafo(paragrafo, substituicoes)
        for paragrafo in secao.footer.paragraphs:
            _substituir_paragrafo(paragrafo, substituicoes)


def _substituir_paragrafo(paragrafo: Paragraph, substituicoes: dict[str, str]) -> None:
    texto = paragrafo.text
    if not texto:
        return
    alterou = False
    for antigo, novo in substituicoes.items():
        if antigo in texto:
            texto = texto.replace(antigo, novo)
            alterou = True
    if alterou:
        if paragrafo.runs:
            paragrafo.runs[0].text = texto
            for run in paragrafo.runs[1:]:
                run.text = ""
        else:
            paragrafo.text = texto


def definir_celula(celula, texto: str) -> None:
    if celula.paragraphs:
        paragrafo = celula.paragraphs[0]
        if paragrafo.runs:
            paragrafo.runs[0].text = texto
            for run in paragrafo.runs[1:]:
                run.text = ""
        else:
            paragrafo.text = texto
    else:
        celula.text = texto


def reconstruir_tabela_pacotes(tabela: Table, grupos: list[dict[str, Any]], total_pf: float) -> None:
    if len(tabela.rows) < 3:
        raise ValueError("Tabela de pacotes deve ter pelo menos cabeçalho, grupo e item modelo.")

    linha_grupo = deepcopy(tabela.rows[1]._tr)
    linha_item = deepcopy(tabela.rows[2]._tr)

    while len(tabela.rows) > 1:
        tabela._tbl.remove(tabela.rows[-1]._tr)

    def append_row(modelo_tr, preenchimento) -> None:
        tabela._tbl.append(deepcopy(modelo_tr))
        row = tabela.rows[-1]
        preenchimento(row)

    for idx, grupo in enumerate(grupos, start=1):
        gid = grupo.get("id") or f"F{idx}"

        def preencher_grupo(row, gid=gid, grupo=grupo):
            definir_celula(row.cells[0], gid)
            definir_celula(row.cells[1], grupo["nome"])
            definir_celula(row.cells[2], "")

        append_row(linha_grupo, preencher_grupo)

        for item in grupo.get("itens", []):
            def preencher_item(row, item=item):
                nome = item["nome"]
                tipo = item.get("tipo", "")
                if tipo and f" - {tipo}" not in nome:
                    nome = f"{nome} - {tipo}"
                pf = item.get("pf_local") or item.get("pf") or 0
                pf_txt = f"{float(pf):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                definir_celula(row.cells[0], "")
                definir_celula(row.cells[1], nome)
                definir_celula(row.cells[2], pf_txt)

            append_row(linha_item, preencher_item)

    total_txt = f"{float(total_pf):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    def preencher_total(row):
        definir_celula(row.cells[0], "")
        definir_celula(row.cells[1], "TOTAL DE PONTOS DE FUNÇÃO")
        definir_celula(row.cells[2], total_txt)

    append_row(linha_item, preencher_total)
