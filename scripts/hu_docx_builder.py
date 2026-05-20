"""Gera DOCX de HU a partir do template Word."""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from docx_utils import definir_celula
from hu_parser import HistoriaUsuario


def gerar_docx_hu(
    hu: HistoriaUsuario,
    template: Path,
    saida: Path,
    config: dict[str, Any],
) -> None:
    saida.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, saida)
    doc = Document(str(saida))

    projeto = hu.projeto or config.get("nome_projeto", "")
    _definir_titulo(doc, f"{hu.codigo} – {hu.titulo.upper()}", f"PROJETO: {projeto}")
    _atualizar_historico(doc.tables[0], hu.historico, config)
    _reconstruir_corpo(doc, hu, config)
    _atualizar_aprovacao(doc, hu, config)

    doc.save(str(saida))


def _definir_titulo(doc: Document, titulo: str, projeto: str) -> None:
    for para in doc.paragraphs:
        if "HU.XX" in para.text:
            _definir_texto_paragrafo(para, titulo)
        elif para.text.startswith("PROJETO:"):
            _definir_texto_paragrafo(para, projeto)


def _definir_texto_paragrafo(para: Paragraph, texto: str) -> None:
    if para.runs:
        para.runs[0].text = texto
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.text = texto


def _atualizar_historico(tabela, historico, config: dict[str, Any]) -> None:
    while len(tabela.rows) > 2:
        tabela._tbl.remove(tabela.rows[-1]._tr)

    itens = historico or [
        type("H", (), {
            "data": "",
            "resumo": "Criação do documento",
            "versao": "1.0",
            "responsavel": config.get("empresa", ""),
        })()
    ]

    for item in itens:
        row = tabela.add_row()
        definir_celula(row.cells[0], item.data)
        definir_celula(row.cells[1], item.resumo)
        definir_celula(row.cells[2], item.versao)
        definir_celula(row.cells[3], item.responsavel)


def _reconstruir_corpo(doc: Document, hu: HistoriaUsuario, config: dict[str, Any]) -> None:
    inicio_para = _paragrafo_por_texto(doc, "VISÃO GERAL")
    fim_para = _paragrafo_por_texto(doc, "APROVAÇÃO DO REQUISITO", exato=True)
    if inicio_para is None or fim_para is None:
        raise ValueError("Template HU inválido: seções VISÃO GERAL ou APROVAÇÃO não encontradas.")

    _remover_entre_paragrafos(inicio_para, fim_para)
    _definir_texto_paragrafo(inicio_para, "")

    inserir = _make_inserter(doc, inicio_para)

    inserir("normal", "VISÃO GERAL")
    inserir("normal", hu.visao_geral or "")
    inserir("normal", "")
    inserir("normal", "SUMÁRIO")
    for item in hu.sumario:
        inserir("normal", f"• {item}")
    if not hu.sumario:
        inserir("normal", "")
    inserir("normal", "")

    epico_titulo = hu.epico_titulo or "Épico"
    if not epico_titulo.upper().startswith("ÉPICO"):
        epico_titulo = f"ÉPICO: {epico_titulo}"
    inserir("Heading 1", epico_titulo.upper() if "ÉPICO:" in epico_titulo.upper() else epico_titulo)
    inserir("normal", hu.epico_texto or "")

    hist_titulo = hu.historia_titulo or f"HISTÓRIA DE USUÁRIO: {hu.titulo.title()}"
    if not hist_titulo.upper().startswith("HISTÓRIA"):
        hist_titulo = f"HISTÓRIA DE USUÁRIO: {hu.titulo.title()}"
    inserir("Heading 1", hist_titulo.upper() if "HISTÓRIA" in hist_titulo.upper() else hist_titulo)
    inserir("normal", hu.historia_texto or "")

    inserir("Heading 1", "PRÉ-REQUISITOS")
    if hu.pre_requisitos:
        for item in hu.pre_requisitos:
            inserir("normal", f"• {item}")
    else:
        inserir("normal", "Nenhum")

    inserir("Heading 1", "DEPENDÊNCIAS")
    inserir("normal", "Passadas")
    if hu.dependencias_passadas:
        for item in hu.dependencias_passadas:
            inserir("normal", f"• {item}")
    else:
        inserir("normal", "Nenhuma")
    inserir("normal", "Futuras")
    if hu.dependencias_futuras:
        for item in hu.dependencias_futuras:
            inserir("normal", f"• {item}")
    else:
        inserir("normal", "Nenhuma")

    inserir("Heading 1", "INTERFACE DE USUÁRIO")
    for tela in hu.telas:
        inserir("Heading 2", f"TELA {tela.numero} - {tela.titulo.upper()}")
        caminho = tela.caminho_menu or "—"
        inserir("normal", f"Caminho no menu: {caminho}")
        inserir("normal", "")
        if tela.screenshot:
            inserir("normal", f"[Protótipo: {tela.screenshot}]")
        else:
            inserir("normal", "[ESPAÇO PARA PROTÓTIPO]")
        inserir("Heading 3", f"COMPORTAMENTO ESPERADO TELA {tela.numero}")
        for idx, cen in enumerate(tela.cenarios, start=1):
            inserir("normal", f"Cenário {idx}: {cen.titulo}")
            inserir("normal", f"Dado que {cen.dado}" if cen.dado else "Dado que")
            inserir("normal", f"Quando {cen.quando}" if cen.quando else "Quando")
            inserir("normal", f"Então {cen.entao}" if cen.entao else "Então")
        inserir("normal", "")

    inserir("normal", "")


def _atualizar_aprovacao(doc: Document, hu: HistoriaUsuario, config: dict[str, Any]) -> None:
    empresa = config.get("empresa", "Sudoeste Informática")
    textos = hu.aprovacao_paragrafos or [
        t.format(empresa=empresa) for t in config.get("texto_aprovacao", [])
    ]

    idx_para = _paragrafo_por_texto(doc, "APROVAÇÃO DO REQUISITO", exato=True)
    if idx_para is None:
        return

    bloco = []
    capturando = False
    for para in doc.paragraphs:
        if para._p is idx_para._p:
            capturando = True
            continue
        if capturando:
            if para.text.strip() and not para.text.strip().startswith("|"):
                bloco.append(para)
            if len(bloco) >= 3:
                break
            if para.style.name.startswith("Heading") and bloco:
                break

    for i, texto in enumerate(textos):
        if i < len(bloco):
            _definir_texto_paragrafo(bloco[i], texto)

    if len(doc.tables) >= 2:
        tabela = doc.tables[1]
        while len(tabela.rows) > 2:
            tabela._tbl.remove(tabela.rows[-1]._tr)
        assinaturas = hu.assinaturas or [{}] * 3
        for assin in assinaturas:
            row = tabela.add_row()
            definir_celula(row.cells[0], assin.get("data", ""))
            definir_celula(row.cells[1], assin.get("setor", ""))
            definir_celula(row.cells[2], assin.get("nome", ""))
            definir_celula(row.cells[3], assin.get("assinatura", ""))


def _paragrafo_por_texto(doc: Document, texto_parcial: str, exato: bool = False) -> Paragraph | None:
    alvo = texto_parcial.upper()
    for para in doc.paragraphs:
        texto = para.text.strip().upper()
        if exato:
            if texto == alvo:
                return para
        elif alvo in texto:
            return para
    return None


def _remover_entre_paragrafos(inicio: Paragraph, fim: Paragraph) -> None:
    while True:
        proximo = inicio._p.getnext()
        if proximo is None or proximo is fim._p:
            break
        parent = proximo.getparent()
        if parent is not None:
            parent.remove(proximo)


def _indice_paragrafo(doc: Document, texto_parcial: str) -> int | None:
    para = _paragrafo_por_texto(doc, texto_parcial)
    if para is None:
        return None
    for i, p in enumerate(doc.paragraphs):
        if p._p is para._p:
            return i
    return None


def _remover_paragrafos_intervalo_legacy(doc: Document, inicio: int, fim: int) -> None:
    body = doc.element.body
    paragrafos = [p for p in body if p.tag.endswith("p")]
    if inicio >= len(paragrafos) or fim >= len(paragrafos):
        return
    for el in paragrafos[inicio:fim]:
        body.remove(el)


def _make_inserter(doc: Document, anchor: Paragraph):
    atual = anchor

    def inserir(estilo: str, texto: str) -> None:
        nonlocal atual
        novo = _inserir_paragrafo_apos(doc, atual, estilo, texto)
        atual = novo

    return inserir


def _inserir_paragrafo_apos(doc: Document, paragrafo: Paragraph, estilo: str, texto: str) -> Paragraph:
    novo_p = OxmlElement("w:p")
    paragrafo._p.addnext(novo_p)
    novo = Paragraph(novo_p, paragrafo._parent)
    try:
        novo.style = estilo
    except KeyError:
        novo.style = "normal"
    if texto:
        novo.add_run(texto)
    return novo
