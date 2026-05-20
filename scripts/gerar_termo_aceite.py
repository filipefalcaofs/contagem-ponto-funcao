#!/usr/bin/env python3
"""Gera Termo de Entrega e Aceite (DOCX fiel + PDF fiel)."""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from typing import Any

from docx import Document

from docx_utils import reconstruir_tabela_pacotes, substituir_texto_documento
from ifpug_calc import calcular_pf
from render_documento import renderizar_arquivo
from gerar_pdf import html_para_pdf, rodape_termo


def formatar_pf(valor: float | int) -> str:
    return f"{float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def calcular_total(dados: dict[str, Any]) -> float:
    termo = dados.get("termo", {})
    if termo.get("total_pf") is not None:
        return float(termo["total_pf"])
    total = 0.0
    for grupo in dados.get("grupos", []):
        for item in grupo.get("itens", []):
            pf = item.get("pf_local") or item.get("pf")
            if pf is None:
                pf = calcular_pf(item.get("tipo", ""), item.get("td"), item.get("ar"))
            total += float(pf or 0)
    return total


def gerar_docx(dados: dict[str, Any], template: Path, saida: Path) -> None:
    total = calcular_total(dados)
    termo = dados.get("termo", {})
    total_fmt = formatar_pf(total)

    saida.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, saida)
    doc = Document(str(saida))

    substituicoes = {
        "1.047,00": total_fmt,
        "contagem-sigvisa-desenvolvimento.xlsx": termo.get("planilha_anexa", "contagem-desenvolvimento.xlsx"),
        "05 de maio de 2026": termo.get("data", "").replace("Salvador-BA, ", ""),
        "001/2023": termo.get("processo_licitatorio", "001/2023"),
        "255/2023": termo.get("contrato", "255/2023"),
        "77751/2023": termo.get("processo_administrativo", "77751/2023"),
    }
    substituir_texto_documento(doc, substituicoes)

    if len(doc.tables) >= 2:
        reconstruir_tabela_pacotes(doc.tables[1], dados.get("grupos", []), total)

    doc.save(str(saida))


def gerar_pdf(dados_json: Path, base_dir: Path, saida_pdf: Path) -> None:
    html_temp = saida_pdf.with_suffix(".html")
    renderizar_arquivo("termo-entrega-aceite.html", dados_json, html_temp, base_dir)
    html_para_pdf(html_temp, saida_pdf, rodape=rodape_termo())
    html_temp.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera Termo de Entrega e Aceite")
    parser.add_argument("json", type=Path)
    parser.add_argument("-o", "--saida", type=Path, help="DOCX de saída")
    parser.add_argument("--pdf", type=Path, help="PDF de saída (layout fiel)")
    parser.add_argument(
        "-t",
        "--template",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "templates" / "termo-entrega-aceite-modelo.docx",
    )
    args = parser.parse_args()
    base = Path(__file__).resolve().parent.parent

    with args.json.open(encoding="utf-8") as f:
        dados = json.load(f)

    if args.saida:
        gerar_docx(dados, args.template, args.saida)
        print(f"Termo DOCX: {args.saida}")

    if args.pdf:
        gerar_pdf(args.json, base, args.pdf)
        print(f"Termo PDF: {args.pdf}")

    if not args.saida and not args.pdf:
        docx = args.json.with_suffix(".docx")
        pdf = args.json.with_suffix(".pdf")
        gerar_docx(dados, args.template, docx)
        gerar_pdf(args.json, base, pdf)
        print(f"Termo DOCX: {docx}")
        print(f"Termo PDF: {pdf}")


if __name__ == "__main__":
    main()
