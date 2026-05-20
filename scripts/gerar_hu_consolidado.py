#!/usr/bin/env python3
"""Consolida todas as HUs em um único DOCX."""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

from hu_config import carregar_config, listar_hus_md, pasta_hus
from hu_docx_builder import gerar_docx_hu
from hu_parser import parsear_arquivo


def gerar_consolidado(
    arquivos_md: list[Path],
    template: Path,
    destino: Path,
    config: dict,
    projeto: Path,
) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    temp_dir = destino.parent / ".hu-temp-consolidado"
    temp_dir.mkdir(exist_ok=True)

    partes: list[Path] = []
    for md in arquivos_md:
        hu = parsear_arquivo(md, projeto_padrao=config.get("nome_projeto", ""))
        parcial = temp_dir / f"{hu.codigo}.docx"
        gerar_docx_hu(hu, template, parcial, config)
        partes.append(parcial)

    if not partes:
        raise ValueError("Nenhuma HU para consolidar.")

    shutil.copy2(partes[0], destino)
    master = Document(str(destino))

    for parcial in partes[1:]:
        master.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc = Document(str(parcial))
        for element in doc.element.body:
            if element.tag.endswith("sectPr"):
                continue
            master.element.body.append(element)

    master.save(str(destino))

    for p in partes:
        p.unlink(missing_ok=True)
    if temp_dir.exists() and not any(temp_dir.iterdir()):
        temp_dir.rmdir()


def main() -> None:
    parser = argparse.ArgumentParser(description="Consolida HUs em um DOCX")
    parser.add_argument("--projeto", type=Path, default=Path.cwd())
    parser.add_argument("--config", type=Path)
    parser.add_argument("-o", "--saida", type=Path)
    parser.add_argument(
        "-t",
        "--template",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "templates" / "hu-modelo.docx",
    )
    args = parser.parse_args()

    projeto = args.projeto.resolve()
    config = carregar_config(args.config, projeto)
    pasta = pasta_hus(projeto, config)
    arquivos = listar_hus_md(pasta)
    if not arquivos:
        print(f"Nenhuma HU em {pasta}", file=sys.stderr)
        sys.exit(1)

    nome_proj = config.get("nome_projeto", projeto.name).split("-")[0].strip().replace(" ", "-")
    saida = args.saida or (pasta / f"HUs-{nome_proj.upper()}-CONSOLIDADO.docx")
    gerar_consolidado(arquivos, args.template, saida, config, projeto)
    print(f"Consolidado: {saida}")


if __name__ == "__main__":
    main()
